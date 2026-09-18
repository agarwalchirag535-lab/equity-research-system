#!/usr/bin/env python3
"""
Independent verifier for BSE_Financial_Model.xlsx.

Checks:
  A. Recalculation with `formulas` engine; assert no Excel error cells.
  B. Balance sheet balances — projected balance-check row ~ 0.
  C. DCF sanity — WACC>g, EV>0, value/share>0 finite; IS revenue ties to actual.
  D. No hardcoded computed numbers in projected columns of formula tabs.
  E. Key projected formulas reference Assumptions.
  F. Formatting doctrine (blue input / green cross-sheet / black same-sheet).
"""
import sys, json, math
import openpyxl

XLSX = "BSE_Financial_Model.xlsx"
ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")
BLUE = "FF0000FF"; GREEN = "FF008000"; BLACK = "FF000000"

results = []
def rec(name, ok, detail=""):
    results.append((name, ok, detail))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))


def recalc():
    try:
        import formulas, logging
        logging.getLogger("formulas").setLevel(logging.ERROR)
    except Exception as e:
        rec("A. recalc (formulas engine)", False, f"import failed: {e}")
        return {}
    xl = formulas.ExcelModel().loads(XLSX).finish()
    sol = xl.calculate()
    errcells, values = [], {}
    for k, v in sol.items():
        try:
            val = v.value[0, 0] if hasattr(v.value, "shape") else v.value
        except Exception:
            val = getattr(v, "value", None)
        values[k] = val
        s = str(val)
        if any(e in s for e in ERRORS):
            errcells.append((k, s))
    rec("A. recalc — no Excel error cells", len(errcells) == 0,
        "" if not errcells else f"{len(errcells)} err e.g. {errcells[:5]}")
    return values


def find(values, sheet, cell):
    tail = f"{sheet.upper()}'!{cell.upper()}"
    for k, v in values.items():
        ku = k.upper()
        if ku.endswith(tail) or ku.endswith(f"{sheet.upper()}!{cell.upper()}"):
            return v
    return None


def main():
    rm = json.load(open("row_maps.json"))
    BS, DCF, IS, SEG, A = rm["BS_ROW"], rm["DCF_ROW"], rm["IS_ROW"], rm["SEG_ROW"], rm["A_ROW"]
    values = recalc()

    if values:
        chk = BS["check"]
        bad = []
        for col in ("E", "F", "G"):
            v = find(values, "Balance Sheet", f"{col}{chk}")
            if v is None or abs(float(v)) > 0.5:
                bad.append((col, v))
        rec("B. Balance sheet balances (projected ~0)", len(bad) == 0,
            "" if not bad else f"non-zero: {bad}")

        wacc = find(values, "DCF Valuation", f"E{DCF['wacc']}")
        tg = find(values, "DCF Valuation", f"E{DCF['tg']}")
        ev = find(values, "DCF Valuation", f"E{DCF['ev']}")
        vps = find(values, "DCF Valuation", f"E{DCF['value_per_share']}")
        try:
            ok = bool(wacc > tg and ev > 0 and vps > 0 and math.isfinite(vps))
        except Exception:
            ok = False
        rec("C. DCF sanity (WACC>g, EV>0, vps>0)", ok,
            (f"WACC={wacc:.4f} g={tg:.4f} EV={ev:,.0f}cr vps=Rs{vps:,.1f}" if ok
             else f"WACC={wacc} g={tg} EV={ev} vps={vps}"))

        rev26 = find(values, "Income Statement", f"D{IS['revenue']}")
        rec("C2. IS revenue FY26 ties (~4834 cr)",
            rev26 is not None and abs(float(rev26) - 4833.95) < 5.0,
            f"FY26 revenue = {rev26}")

    wb = openpyxl.load_workbook(XLSX)
    offenders = []
    for tab in ["Income Statement", "Balance Sheet", "DCF Valuation"]:
        ws = wb[tab]
        for col in ("E", "F", "G"):
            for row in range(6, ws.max_row + 1):
                c = ws[f"{col}{row}"]
                if isinstance(c.value, (int, float)):
                    offenders.append(f"{tab}!{col}{row}={c.value}")
    rec("D. No hardcoded numbers in projected cols of formula tabs",
        len(offenders) == 0, "" if not offenders else f"{len(offenders)} e.g. {offenders[:6]}")

    # -------- E. ALL projected formulas in IS/BS/Segment trace to Assumptions --------
    # Build a same-workbook dependency graph and check each projected E/F/G formula cell
    # reaches an 'Assumptions!' reference transitively. Documented allowed exceptions:
    #   - Balance Sheet opening net block anchored to prior actual (E of ppe_open)
    #   - Segment consolidation scalar held flat from historical ratio (consol_scalar row)
    import re as _re
    allowed_rows = {
        ("Balance Sheet", BS["ppe_open"]),
        ("Segment Revenue Build", SEG["consol_scalar"]),
    }
    cellref = _re.compile(r"(?:'([^']+)'|([A-Za-z ]+))?!?\$?([A-Z]{1,2})\$?(\d+)")

    def traces_to_assumptions(tab, cell, seen=None, depth=0):
        if seen is None:
            seen = set()
        key = (tab, cell)
        if key in seen or depth > 40:
            return False
        seen.add(key)
        f = wb[tab][cell].value
        if not isinstance(f, str) or not f.startswith("="):
            return False
        if "Assumptions!" in f:
            return True
        # walk same-tab and cross-tab references
        for m in _re.finditer(r"(?:'([^']+)'!)?\$?([A-Z]{1,2})\$?(\d+)", f):
            rtab = m.group(1) or tab
            rcell = f"{m.group(2)}{m.group(3)}"
            if rtab == "Assumptions":
                return True
            if rtab in wb.sheetnames and traces_to_assumptions(rtab, rcell, seen, depth + 1):
                return True
        return False

    untraced = []
    for tab in ["Income Statement", "Balance Sheet", "Segment Revenue Build"]:
        ws = wb[tab]
        for col in ("E", "F", "G"):
            for row in range(6, ws.max_row + 1):
                c = ws[f"{col}{row}"]
                if isinstance(c.value, str) and c.value.startswith("="):
                    if (tab, row) in allowed_rows:
                        continue
                    if not traces_to_assumptions(tab, f"{col}{row}"):
                        untraced.append(f"{tab}!{col}{row}={c.value}")
    rec("E. ALL projected formulas trace to Assumptions (ex documented anchors)",
        len(untraced) == 0, "" if not untraced else f"{len(untraced)}: {untraced[:8]}")

    def fc(ws, cell):
        col = ws[cell].font.color
        return col.rgb if col and col.rgb else None
    issues = []
    ws_is = wb["Income Statement"]
    if fc(ws_is, f"D{IS['revenue']}") != GREEN:
        issues.append("IS revenue not green")
    if fc(ws_is, f"E{IS['total_income']}") not in (BLACK, None):
        issues.append("IS total_income not black")
    wsA = wb["Assumptions"]
    if fc(wsA, f"E{A['tax_rate']}") != BLUE:
        issues.append("tax input not blue")
    if wsA[f"E{A['tax_rate']}"].fill.fgColor.rgb != "FFFFFF00":
        issues.append("tax input not yellow")
    rec("F. Formatting doctrine (sampled)", len(issues) == 0,
        "" if not issues else str(issues))

    npass = sum(1 for _, ok, _ in results if ok)
    print(f"\n{'='*60}\n{npass}/{len(results)} checks passed")
    return 0 if npass == len(results) else 1


if __name__ == "__main__":
    sys.exit(main())

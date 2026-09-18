#!/usr/bin/env python3
"""Independent verifier for BSE_Financial_Model_v2.xlsx."""
import sys, json, math, re
import openpyxl

XLSX = "BSE_Financial_Model_v2.xlsx"
ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NUM!", "#NULL!")
BLUE, GREEN, BLACK = "FF0000FF", "FF008000", "FF000000"
FORMULA_TABS = ["Income Statement", "Balance Sheet", "Segment Revenue Build", "Cash Flow",
                "DCF Valuation", "Comps", "Sensitivity", "Checks", "Drivers & History"]

res = []
def rec(n, ok, d=""):
    res.append((n, ok, d))
    print(f"[{'PASS' if ok else 'FAIL'}] {n}" + (f" — {d}" if d else ""))


def recalc():
    import formulas, logging
    logging.getLogger("formulas").setLevel(logging.ERROR)
    sol = formulas.ExcelModel().loads(XLSX).finish().calculate()
    vals, errs = {}, []
    for k, v in sol.items():
        try:
            x = v.value[0, 0] if hasattr(v.value, "shape") else v.value
        except Exception:
            x = getattr(v, "value", None)
        vals[k] = x
        if any(e in str(x) for e in ERRORS):
            errs.append((k, str(x)))
    rec("A. recalc — no Excel error cells", not errs,
        "" if not errs else f"{len(errs)} e.g. {errs[:4]}")
    return vals


def get(v, sheet, cell):
    t = f"{sheet.upper()}'!{cell.upper()}"
    for k, x in v.items():
        if k.upper().endswith(t) or k.upper().endswith(f"{sheet.upper()}!{cell.upper()}"):
            return x
    return None


def main():
    rm = json.load(open("row_maps_v2.json"))
    AR, DR, SG, IS, BS, CF, DC, CK = (rm[k] for k in ("AR", "DR", "SG", "IS", "BS", "CF", "DC", "CK"))
    v = recalc()
    wb = openpyxl.load_workbook(XLSX)

    # ---- B. all Checks-tab statuses ----
    ws = wb["Checks"]
    fails = []
    for row in range(6, ws.max_row + 1):
        st = get(v, "Checks", f"E{row}")
        if isinstance(st, str) and st.strip() == "REVIEW":
            fails.append((ws[f"A{row}"].value, get(v, "Checks", f"B{row}"), get(v, "Checks", f"C{row}")))
    rec("B. Checks tab — every reconciliation PASSes", not fails,
        "" if not fails else f"{len(fails)} REVIEW: {fails[:3]}")

    # ---- C. balance sheet balances ----
    bad = []
    for c in "EFG":
        x = get(v, "Balance Sheet", f"{c}{BS['check']}")
        if x is None or abs(float(x)) > 0.5:      # note: must not use `or` on a 0.0 value
            bad.append((c, x))
    rec("C. Balance sheet balances in all projected years", not bad, str(bad) if bad else "all = 0.00")

    # ---- D. DCF integrity: no double count, WACC>g, coherent reverse DCF ----
    wacc = float(get(v, "DCF Valuation", f"E{DC['wacc']}"))
    g = float(get(v, "DCF Valuation", f"E{DC['g']}"))
    ev = float(get(v, "DCF Valuation", f"E{DC['ev_ops']}"))
    vps = float(get(v, "DCF Valuation", f"E{DC['vps']}"))
    ig = float(get(v, "DCF Valuation", f"E{DC['implied_g']}"))
    ggap = float(get(v, "DCF Valuation", f"E{DC['g_gap']}"))
    ok = wacc > g and ev > 0 and vps > 0 and math.isfinite(vps) and ggap < 0
    rec("D. DCF integrity (WACC>g, EV>0, reverse-DCF coherent)", ok,
        f"WACC={wacc:.2%} g={g:.2%} EV={ev:,.0f} vps=Rs{vps:,.0f} implied_g={ig:.2%} (g-WACC={ggap:.2%})")

    # ---- E. the v1 double-count is actually gone ----
    # operating EBIT must NOT include investment income: check the IS formula
    f_ebit = wb["Income Statement"][f"E{IS['ebit_op']}"].value
    no_inv = str(IS['inv_income']) not in re.findall(r"[A-G](\d+)", f_ebit or "")
    # and the DCF must add the treasury exactly once
    f_eq = wb["DCF Valuation"][f"E{DC['eq_val']}"].value
    once = f_eq.count(f"E{DC['treasury']}") == 1
    rec("E. v1 double-count removed (EBIT excludes treasury income; corpus added once)",
        no_inv and once, f"ebit_op='{f_ebit}' ; eq_val='{f_eq}'")

    # ---- F. three-driver architecture is real ----
    f_levy = wb["Income Statement"][f"E{IS['sebi_levy']}"].value or ""
    f_clear = wb["Income Statement"][f"E{IS['clearing']}"].value or ""
    f_deriv = wb["Segment Revenue Build"][f"E{SG['deriv']}"].value or ""
    arch = (f"{DR['notional_adt']}" in f_levy and "Drivers" in f_levy
            and f"{DR['contracts']}" in f_clear and "Drivers" in f_clear
            and f"{DR['premium_adt']}" in f_deriv and "Drivers" in f_deriv)
    rec("F. Three-driver architecture (levy<-notional, clearing<-contracts, revenue<-premium)",
        arch, f"levy='{f_levy[:52]}' clearing='{f_clear[:46]}'")

    # ---- G. wedge drives notional ----
    f_not = wb["Drivers & History"][f"E{DR['notional_adt']}"].value or ""
    rec("G. Notional is DERIVED from premium / wedge (wedge is live)",
        f"{DR['premium_adt']}" in f_not and f"{DR['wedge']}" in f_not, f"'{f_not}'")

    # ---- H. no hardcoded numbers in projected columns of formula tabs ----
    off = []
    for tab in ["Income Statement", "Balance Sheet", "DCF Valuation", "Cash Flow"]:
        w = wb[tab]
        for c in "EFG":
            for row in range(6, w.max_row + 1):
                cell = w[f"{c}{row}"]
                if isinstance(cell.value, (int, float)):
                    # DCF/Comps single-input col E is legitimate only if blue-font input
                    rgb = cell.font.color.rgb if cell.font.color else None
                    if rgb != BLUE:
                        off.append(f"{tab}!{c}{row}={cell.value}")
    rec("H. No non-blue hardcoded numbers in projected columns", not off,
        "" if not off else f"{len(off)} e.g. {off[:5]}")

    # ---- I. transitive trace to Assumptions ----
    allowed = {("Balance Sheet", BS["nb_open"]), ("Segment Revenue Build", SG["scalar"]),
               ("Cash Flow", CF.get("drecv", -1)), ("Cash Flow", CF.get("dfloat", -1))}
    def traces(tab, cell, seen=None, d=0):
        if seen is None: seen = set()
        if (tab, cell) in seen or d > 45: return False
        seen.add((tab, cell))
        f = wb[tab][cell].value
        if not isinstance(f, str) or not f.startswith("="): return False
        if "Assumptions!" in f: return True
        for m in re.finditer(r"(?:'([^']+)'!)?\$?([A-Z]{1,2})\$?(\d+)", f):
            rt = m.group(1) or tab
            if rt == "Assumptions": return True
            if rt in wb.sheetnames and traces(rt, f"{m.group(2)}{m.group(3)}", seen, d + 1):
                return True
        return False
    untraced = []
    for tab in ["Income Statement", "Balance Sheet", "Segment Revenue Build", "Drivers & History"]:
        w = wb[tab]
        for c in "EFG":
            for row in range(6, w.max_row + 1):
                val = w[f"{c}{row}"].value
                if isinstance(val, str) and val.startswith("=") and (tab, row) not in allowed:
                    if not traces(tab, f"{c}{row}"):
                        untraced.append(f"{tab}!{c}{row}={val[:44]}")
    rec("I. All projected formulas trace to Assumptions (ex documented anchors)",
        not untraced, "" if not untraced else f"{len(untraced)}: {untraced[:6]}")

    # ---- J. formatting doctrine ----
    def fc(tab, cell):
        col = wb[tab][cell].font.color
        return col.rgb if col and col.rgb else None
    iss = []
    if fc("Income Statement", f"D{IS['revenue']}") != GREEN: iss.append("IS revenue not green")
    if fc("Income Statement", f"E{IS['total_income']}") not in (BLACK, None): iss.append("total income not black")
    if fc("Assumptions", f"E{AR['tax_rate']}") != BLUE: iss.append("tax input not blue")
    if wb["Assumptions"][f"E{AR['tax_rate']}"].fill.fgColor.rgb != "FFFFFF00": iss.append("tax not yellow")
    if fc("Assumptions", f"H{AR['tax_rate']}") not in (BLACK, None): iss.append("ACTIVE not black")
    rec("J. Formatting doctrine (blue input / green cross / black same)", not iss, str(iss) if iss else "")

    # ---- K. scenario switch is wired ----
    f_act = wb["Assumptions"][f"H{AR['tax_rate']}"].value or ""
    rec("K. Scenario switch wired through the ACTIVE column",
        "CHOOSE" in f_act and "$C$4" in f_act, f"'{f_act}'")

    n = sum(1 for _, o, _ in res if o)
    print(f"\n{'='*64}\n{n}/{len(res)} checks passed")
    return 0 if n == len(res) else 1


if __name__ == "__main__":
    sys.exit(main())
